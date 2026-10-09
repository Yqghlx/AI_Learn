# -*- coding: utf-8 -*-
"""构建《Coding Agent 实战指南》演示文稿 — 深空科技风 1280x720"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.abspath(__file__))
A = os.path.join(ROOT, "assets")
OUT = os.path.join(ROOT, "Coding Agent 实战指南.pptx")

BG = RGBColor(0x0F, 0x1B, 0x2D)
TEAL = RGBColor(0x35, 0xD0, 0xBA)
BLUE = RGBColor(0x4C, 0x8D, 0xFF)
TXT = RGBColor(0xE8, 0xED, 0xF5)
CARD = RGBColor(0x1C, 0x28, 0x39)      # rgba(232,237,245,.06)
CARD_BR = RGBColor(0x29, 0x34, 0x45)   # rgba(232,237,245,.12)
MUT55 = RGBColor(0x8A, 0x94, 0xA6)
MUT50 = RGBColor(0x7B, 0x84, 0x94)
MUT65 = RGBColor(0xA6, 0xAD, 0xB8)
MUT68 = RGBColor(0xA9, 0xB0, 0xBB)
MUT72 = RGBColor(0xAB, 0xB2, 0xBD)
MUT78 = RGBColor(0xBC, 0xC3, 0xCD)
MUT85 = RGBColor(0xC8, 0xCE, 0xD6)
TEAL_BG = RGBColor(0x13, 0x2B, 0x3A)
TEAL_BG2 = RGBColor(0x14, 0x31, 0x40)
TEAL_BR = RGBColor(0x20, 0x6D, 0x6D)
BLUE_BG = RGBColor(0x18, 0x2C, 0x4D)
BLUE_BR = RGBColor(0x2D, 0x54, 0x96)
CODE_BG = RGBColor(0x0C, 0x16, 0x26)

CN = "Microsoft YaHei"
MONO = "Consolas"

def emu(px):
    return Emu(int(px * 9525))

def set_run(r, text, size, color=TXT, bold=False, font=CN, italic=False, spacing=None):
    r.text = text
    r.font.size = Pt(size * 0.75)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", CN if font == CN else font)
    if spacing:
        rPr.set("spc", str(int(spacing * 100)))

def tb(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0, space_after=0):
    """runs: list of paragraphs; each paragraph = list of (text, kwargs) tuples"""
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if space_after:
            p.space_after = Pt(space_after * 0.75)
        for text, kw in para:
            set_run(p.add_run(), text, **kw)
    return box

def rect(slide, x, y, w, h, fill=None, line=None, radius=None, line_w=1.0, shadow=False):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp_type, emu(x), emu(y), emu(w), emu(h))
    if radius is not None:
        try:
            sp.adjustments[0] = radius / min(w, h) if min(w, h) > 0 else 0.1
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    sp.text_frame.paragraphs[0].text = ""
    return sp

def hline(slide, x, y, w, color, thickness=2):
    return rect(slide, x, y, w, thickness, fill=color)

def vline(slide, x, y, h, color, thickness=2):
    return rect(slide, x, y, thickness, h, fill=color)

def arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.0):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    conn.shadow.inherit = False
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn

def pic(slide, path, x, y, w, h):
    return slide.shapes.add_picture(path, emu(x), emu(y), emu(w), emu(h))

def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s

def masked(src, dst, l_a, r_a, bottom=None):
    im = Image.open(src).convert("RGBA")
    w, h = im.size
    overlay = Image.new("RGBA", (w, h), (15, 27, 45, 0))
    px = overlay.load()
    for xx in range(w):
        t = xx / (w - 1)
        a = int(l_a + (r_a - l_a) * t)
        for yy in range(h):
            px[xx, yy] = (15, 27, 45, a)
    if bottom:
        bh = int(h * bottom[0])
        for yy in range(bh, h):
            t = (yy - bh) / (h - bh)
            a = int(bottom[1] * t)
            for xx in range(w):
                r0, g0, b0, a0 = px[xx, yy]
                a1 = min(255, a0 + a)
                px[xx, yy] = (15, 27, 45, a1)
    im = Image.alpha_composite(im, overlay).convert("RGB")
    im.save(dst)

def header(slide, title, sub, accent=None):
    tb(slide, 64, 26, 1000, 46, [[(title, dict(size=36, bold=True))]])
    hline(slide, 64, 82, 56, accent or TEAL, 3)
    tb(slide, 136, 72, 1000, 26, [[(sub, dict(size=16, color=MUT55))]])

def footer(slide, chapter, page):
    rect(slide, 64, 688, 12, 12, fill=RGBColor(0x14, 0x31, 0x40), line=TEAL_BR, radius=2)
    tb(slide, 88, 684, 700, 22, [[("Coding Agent 实战指南 · " + chapter, dict(size=14, color=MUT50))]])
    tb(slide, 1000, 684, 216, 22, [[(page, dict(size=14, color=TEAL, font=MONO))]], align=PP_ALIGN.RIGHT)

def judge_card(slide, x, y, w, h, text, size=17.5):
    rect(slide, x, y, w, h, fill=TEAL_BG, line=TEAL_BR, radius=10)
    rect(slide, x + 24, y + 20, 10, 10, fill=TEAL, radius=5)
    tb(slide, x + 44, y + 14, 120, 22, [[("本课判断", dict(size=16, bold=True, color=TEAL))]])
    tb(slide, x + 24, y + 44, w - 48, h - 58, [[(text, dict(size=size, color=TXT))]], line_spacing=1.6)

def chip(slide, x, y, w, text, color=TXT):
    rect(slide, x, y, w, 30, fill=None, line=CARD_BR, radius=15)
    tb(slide, x, y + 5, w, 20, [[(text, dict(size=14, color=color))]], align=PP_ALIGN.CENTER)

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)

# ---------- P01 封面 ----------
s = new_slide(prs)
masked(os.path.join(A, "cover_agent.png"), os.path.join(A, "p_cover.png"), 240, 40, bottom=(0.45, 200))
pic(s, os.path.join(A, "p_cover.png"), 0, 0, 1280, 720)
hline(s, 64, 158, 96, TEAL, 4)
tb(s, 176, 146, 500, 24, [[("RESEARCH EFFECTIVENESS", dict(size=14, color=TEAL, font=MONO, spacing=3))]])
tb(s, 64, 186, 700, 100, [[("Coding Agent", dict(size=72, bold=True, font=MONO))]])
tb(s, 64, 286, 700, 92, [[("实战指南", dict(size=66, bold=True, color=TEAL))]])
tb(s, 64, 400, 640, 40, [[("国内编码智能体横向对比 × 工程落地最佳实践", dict(size=22, color=MUT72))]])
rect(s, 64, 630, 150, 34, fill=CARD, line=CARD_BR, radius=17)
rect(s, 78, 643, 8, 8, fill=TEAL, radius=4)
tb(s, 96, 637, 120, 22, [[("研发效能内训", dict(size=14, color=TXT))]])
tb(s, 230, 637, 120, 22, [[("2026.10", dict(size=14, color=MUT55, font=MONO))]])
rect(s, 980, 650, 300, 70, fill=BG)
tb(s, 1000, 676, 260, 22, [[("AGENTIC CODING 101", dict(size=13, color=MUT50, font=MONO, spacing=2))]], align=PP_ALIGN.RIGHT)

# ---------- P02 目录 ----------
s = new_slide(prs)
tb(s, 64, 30, 400, 46, [[("目录", dict(size=36, bold=True))]])
hline(s, 64, 82, 56, TEAL, 3)
rect(s, 64, 160, 6, 120, fill=TEAL)
tb(s, 64, 300, 300, 80, [[("四个问题，", dict(size=24, bold=True))], [("一条落地路径", dict(size=24, bold=True))]], line_spacing=1.4)
tb(s, 64, 392, 300, 120, [[("从认知到工具，从方法到团队——每一章回答一个你下周就会遇到的问题。", dict(size=16, color=MUT55))]], line_spacing=1.7)
toc = [
    ("01", "认知篇 · 从补全到自主执行", "Coding Agent 到底和补全工具差在哪？", "P03 - P05"),
    ("02", "格局篇 · 国内 Coding Agent 全景", "四强并立，我的团队该选哪个生态位？", "P06 - P10"),
    ("03", "实践篇 · 上下文、约束与交付循环", "为什么它写的代码总是不对味？", "P11 - P17"),
    ("04", "落地篇 · 团队推广与度量", "如何从个人提效走到团队可复制？", "P18 - P20"),
]
for i, (no, title, q, pg) in enumerate(toc):
    y = 140 + i * 124
    rect(s, 420, y, 796, 110, fill=CARD, line=CARD_BR, radius=10)
    tb(s, 448, y + 30, 70, 46, [[(no, dict(size=30, bold=True, color=TEAL if i == 2 else BLUE, font=MONO))]])
    tb(s, 530, y + 22, 560, 32, [[(title, dict(size=22, bold=True))]])
    tb(s, 530, y + 60, 560, 26, [[(q, dict(size=16, color=MUT55))]])
    tb(s, 1090, y + 44, 110, 24, [[(pg, dict(size=14, color=MUT50, font=MONO))]], align=PP_ALIGN.RIGHT)
footer(s, "目录", "02 / 21")

# ---------- 章节扉页 helper ----------
def section_slide(num, img, title, sub, desc):
    s = new_slide(prs)
    dst = os.path.join(A, "p_ch%s.png" % num)
    masked(img, dst, 240, 90)
    pic(s, dst, 0, 0, 1280, 720)
    tb(s, 880, -30, 400, 340, [[(num, dict(size=260, bold=True, color=RGBColor(0x1A, 0x26, 0x38), font=MONO))]], align=PP_ALIGN.RIGHT)
    tb(s, 64, 210, 400, 26, [[("CHAPTER " + num, dict(size=20, color=TEAL, font=MONO, spacing=4))]])
    hline(s, 250, 226, 72, TEAL, 3)
    tb(s, 64, 250, 700, 90, [[(title, dict(size=64, bold=True, spacing=2))]])
    tb(s, 64, 356, 700, 40, [[(sub, dict(size=24, color=MUT78))]])
    tb(s, 64, 416, 620, 90, [[(desc, dict(size=17, color=MUT55))]], line_spacing=1.7)
    return s

# ---------- P03 ----------
section_slide("01", os.path.join(A, "ch_cognition.png"), "认知篇", "先换心智模型，再谈工具",
              "它不是更快的补全，而是会调用工具、规划步骤的执行者——你如何理解它，决定了你能用它走多远。")

# ---------- P04 ----------
s = new_slide(prs)
header(s, "三次范式跃迁", "每一次跃迁，改变的都是\u201c谁在规划\u201d这件事")
vline(s, 84, 175, 420, CARD_BR, 2)
levels = [
    ("2021 -", "代码补全", "人在写，AI 猜下一个词。Tab 键就是全部交互，出错成本极低。", False),
    ("2023 -", "对话改写", "人规划，AI 生成整段代码。选中 - 提问 - 粘贴，心智仍是\u201c我来改\u201d。", False),
    ("2025 -", "自主执行", "AI 接收目标，自己拆任务、调工具、改文件、跑测试；人定边界、给约束、做验收。", True),
]
for i, (yr, name, desc, hl) in enumerate(levels):
    y = 168 + i * 160
    ring = TEAL_BG2 if hl else BLUE_BG
    edge = None if hl else BLUE
    dot_c = BG if hl else BLUE
    sp = rect(s, 64, y + 44, 40, 40, fill=ring, line=edge, radius=20, line_w=2)
    rect(s, 76 if not hl else 72, y + 56 if not hl else y + 60, 16 if hl else 10, 16 if hl else 10, fill=dot_c, radius=8 if hl else 5)
    rect(s, 122, y, 558, 138, fill=TEAL_BG if hl else CARD, line=TEAL_BR if hl else CARD_BR, radius=10)
    tb(s, 146, y + 26, 120, 26, [[(yr, dict(size=17, color=TEAL if hl else BLUE, font=MONO))]])
    tb(s, 250, y + 20, 200, 32, [[(name, dict(size=22, bold=True))]])
    tb(s, 146, y + 64, 510, 56, [[(desc, dict(size=16, color=MUT68))]], line_spacing=1.5)
tb(s, 748, 152, 468, 120, [[("大多数团队的真实状态：手里拿着第三级的工具，脑子里还是第一级的心智——把它当猜词机器，用完一句\u201c不好用\u201d就关掉了。", dict(size=20, color=TXT))]], line_spacing=1.7)
rect(s, 748, 296, 4, 92, fill=BLUE)
tb(s, 766, 296, 450, 96, [[("工具每升一级，人的职责就后退一步：从\u201c写字\u201d退到\u201c改字\u201d，再退到\u201c定标准、验结果\u201d。", dict(size=17, color=MUT65))]], line_spacing=1.7)
judge_card(s, 748, 470, 468, 160, "心智模型不换，工具永远停在第一级。本课后面所有实践，都建立在第三级的工作方式上。", 18)
footer(s, "认知篇", "04 / 21")

# ---------- P05 ----------
s = new_slide(prs)
header(s, "拆开看：Agent 的五层能力栈", "你抱怨的\u201c不够聪明\u201d，多数时候不是模型层的问题")
stack = [
    ("05", "反馈层", "测试、Lint、CI、人工评审——让 Agent 知道自己做对了没有", False),
    ("04", "记忆层", "跨会话记忆、项目约定沉淀，避免每次从零认识你的仓库", False),
    ("03", "上下文层", "规则文件、任务简报、代码库检索——本课的主战场", True),
    ("02", "工具层", "读写文件、执行终端、检索代码库、MCP 外部服务", False),
    ("01", "模型层", "推理与代码能力底座，如 Qwen3-Coder（SWE-bench Verified 72.3%）", False),
]
for i, (no, name, desc, hl) in enumerate(stack):
    y = 152 + i * 98
    rect(s, 64, y, 640, 86, fill=TEAL_BG if hl else CARD, line=TEAL_BR if hl else CARD_BR, radius=10)
    rect(s, 64, y, 4, 86, fill=TEAL if hl else BLUE)
    tb(s, 92, y + 24, 60, 36, [[(no, dict(size=26, bold=True, color=TEAL if hl else BLUE, font=MONO))]])
    tb(s, 160, y + 26, 120, 32, [[(name, dict(size=21, bold=True))]])
    tb(s, 290, y + 16, 400, 60, [[(desc, dict(size=15.5, color=MUT68))]], line_spacing=1.4, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 748, 152, 468, 130, [[("同一个模型，在不同团队手里产出天差地别——差距几乎都出在第 2、3、5 层：工具接得全不全，上下文给得准不准，反馈回路闭没闭。", dict(size=20, color=TXT))]], line_spacing=1.7)
rect(s, 748, 306, 4, 88, fill=BLUE)
tb(s, 766, 306, 450, 92, [[("模型层决定下限：它保证\u201c能写出能跑的代码\u201d。但决定上限的，是你喂给它的上下文质量。", dict(size=17, color=MUT65))]], line_spacing=1.7)
judge_card(s, 748, 470, 468, 160, "与其纠结换哪个模型，不如先补齐上下文层——这是投入产出比最高的一层，也是第三章的全部内容。", 18)
footer(s, "认知篇", "05 / 21")

# ---------- P06 ----------
section_slide("02", os.path.join(A, "ch_landscape.png"), "格局篇", "四强并立，选的是生态位",
              "国内第一梯队已经贴身缠斗——这一章帮你在 30 分钟内建立选型框架，而不是再刷一遍营销稿。")

# ---------- P07 ----------
s = new_slide(prs)
header(s, "国内格局：四强竞逐，无人通吃", "数据截至 2026 年 5 月 · 公开市场统计口径", TEAL)
rect(s, 1140, -60, 320, 320, fill=RGBColor(0x16, 0x2A, 0x3A), radius=160)
tb(s, 64, 148, 420, 140, [[("28%", dict(size=130, bold=True, color=TEAL, font=MONO))]])
tb(s, 64, 306, 420, 30, [[("腾讯 CodeBuddy 暂居第一", dict(size=20, bold=True))]])
tb(s, 64, 344, 420, 100, [[("但第一名与第四名只差 6 个百分点——这个量级的差距，随时会被下一次版本更新改写。", dict(size=17, color=MUT65))]], line_spacing=1.65)
bars = [
    ("腾讯 CodeBuddy", "增速最快 · 三形态齐全", 28, True),
    ("阿里 Qoder（原通义灵码）", "用户规模最大", 26, False),
    ("字节 Trae", "免费策略 · 增长迅猛", 24, False),
    ("百度文心快码", "企业级能力最强", 22, False),
]
for i, (name, tag, pct, hl) in enumerate(bars):
    y = 158 + i * 66
    tb(s, 540, y, 220, 24, [[(name, dict(size=15.5, bold=True))]])
    tb(s, 540, y + 26, 220, 20, [[(tag, dict(size=12.5, color=MUT50))]])
    rect(s, 776, y + 2, 380, 26, fill=RGBColor(0x18, 0x24, 0x36), radius=4)
    bw = int(380 * pct / 28)
    rect(s, 776, y + 2, bw, 26, fill=TEAL if hl else BLUE, radius=4)
    tb(s, 1168, y + 2, 50, 26, [[("%d%%" % pct, dict(size=19, bold=True, color=TEAL if hl else BLUE, font=MONO))]])
rect(s, 64, 470, 1152, 130, fill=TEAL_BG, line=TEAL_BR, radius=12)
rect(s, 92, 496, 10, 10, fill=TEAL, radius=5)
tb(s, 112, 490, 200, 22, [[("本课判断", dict(size=16, bold=True, color=TEAL))]])
tb(s, 92, 520, 1096, 70, [[("份额接近意味着\u201c没有错误答案\u201d：别再为\u201c谁第一\u201d浪费选型时间，把精力放在\u201c谁和你的技术栈、云生态、合规要求最匹配\u201d——下一页给出四强的完整拆解。", dict(size=19, color=TXT))]], line_spacing=1.6)
footer(s, "格局篇", "07 / 21")

# ---------- P08 ----------
s = new_slide(prs)
header(s, "四强逐个拆解", "不看广告看疗效：差异点、适配场景，以及它们不会写进官网的\u201c代价\u201d")
prods = [
    ("腾讯 CodeBuddy", "腾讯 · 2025.09", "28% 份额 · 插件/IDE/CLI 三形态",
     ["微信生态深度集成：小程序、小游戏、云开发直接生成可提交成品", "Craft 端到端模式 + Figma 设计稿转代码", "SubAgents 子代理与 Typed Memory 跨会话记忆"],
     "微信 / 企微生态、腾讯云团队", "离开腾讯生态，差异化优势收窄", True),
    ("阿里 Qoder", "阿里 · 原通义灵码", "26% 份额 · SWE-bench 72.3%",
     ["底层 Qwen3.8-Max：2.4 万亿参数、100 万上下文", "工程自动感知、记忆学习、专家团协作", "企业版支持私有化，服务一汽、建信金科"],
     "阿里云生态、金融政务高合规", "长周期自主运行仍需人工验收兜底", False),
    ("字节 Trae", "字节跳动 · 2025.01", "24% 份额 · 全功能免费",
     ["IDE + SOLO 双模式：控制感与自动化自由切换", "豆包 1.5-pro、DeepSeek R1/V3 多模型内置", "国内直连零配置，对新手最友好"],
     "个人开发者、学生、MVP 创业", "企业版暂未推出，团队治理能力待补", False),
    ("百度文心快码", "百度 · 2023.07", "百度内部新增代码 45% AI 生成",
     ["3.5S 多智能体自协同：计划/编码/测试智能体分工", "SPEC 模式：过程可回溯、任意节点可干预", "中文理解行业第一，代码库级 RAG + 安全扫描"],
     "中文复杂业务、强数据安全企业", "个人端生态与社区活跃度相对较弱", False),
]
for i, (name, vendor, metric, pts, fit, cost, hl) in enumerate(prods):
    x = 64 + i * 292
    rect(s, x, 148, 276, 470, fill=CARD, line=CARD_BR, radius=10)
    rect(s, x, 148, 276, 4, fill=TEAL if hl else BLUE, radius=2)
    tb(s, x + 20, 172, 240, 30, [[(name, dict(size=19, bold=True))]])
    tb(s, x + 20, 204, 240, 20, [[(vendor, dict(size=12.5, color=MUT50, font=MONO))]])
    rect(s, x + 20, 232, 236, 34, fill=RGBColor(0x18, 0x24, 0x36), radius=6)
    tb(s, x + 30, 240, 220, 22, [[(metric, dict(size=12.5, color=TEAL if hl else BLUE, font=MONO))]])
    for j, p in enumerate(pts):
        yy = 280 + j * 74
        rect(s, x + 20, yy + 8, 5, 5, fill=BLUE, radius=3)
        tb(s, x + 34, yy, 226, 70, [[(p, dict(size=13.5, color=MUT72))]], line_spacing=1.5)
    rect(s, x + 20, 516, 236, 84, fill=RGBColor(0x18, 0x24, 0x36), radius=6)
    tb(s, x + 32, 522, 214, 36, [[("适合：", dict(size=12.5, color=MUT50, bold=True)), (fit, dict(size=12.5, color=MUT55))]], line_spacing=1.4)
    tb(s, x + 32, 556, 214, 40, [[("代价：", dict(size=13, color=TEAL, bold=True)), (cost, dict(size=13, color=MUT85))]], line_spacing=1.4)
footer(s, "格局篇", "08 / 21")

# ---------- P09 ----------
s = new_slide(prs)
header(s, "横向对比总表", "六个维度一次看全 · 数据截至 2026 年中公开信息")
rect(s, 64, 148, 742, 478, fill=RGBColor(0x14, 0x1F, 0x30), line=CARD_BR, radius=10)
rows, cols = 7, 5
tbl_shape = s.shapes.add_table(rows, cols, emu(76), emu(158), emu(718), emu(458))
tbl = tbl_shape.table
widths = [120, 156, 148, 142, 152]
for i, wd in enumerate(widths):
    tbl.columns[i].width = emu(wd)
data = [
    ["维度", "CodeBuddy", "Qoder", "Trae", "文心快码"],
    ["产品形态", "插件+IDE+CLI", "插件+IDE", "独立 IDE", "插件+IDE"],
    ["核心模型", "混元+DeepSeek", "Qwen3.8-Max", "豆包+DeepSeek", "文心大模型"],
    ["智能体模式", "Craft+SubAgents", "工程感知+记忆学习", "SOLO 模式", "多智能体自协同"],
    ["生态绑定", "微信+腾讯云", "阿里云+云效", "字节系（较轻）", "百度智能云"],
    ["企业能力", "私有化（VPC 内）", "私有化 · 9.9 万/年", "暂未推出", "私有化+RAG+安全扫描"],
    ["价格", "个人免费", "团队 99/人/月", "全免费", "专业版 199/月"],
]
hl_cells = {(1, 1), (3, 1), (5, 4)}
for r_i in range(rows):
    tbl.rows[r_i].height = emu(56 if r_i == 0 else 66)
    for c_i in range(cols):
        cell = tbl.cell(r_i, c_i)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = emu(6)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        if r_i == 0:
            set_run(p.add_run(), data[r_i][c_i], 14, TXT if c_i else TXT, bold=True)
            cell.fill.solid(); cell.fill.fore_color.rgb = BLUE_BG
        else:
            hl = (r_i, c_i) in hl_cells
            color = TEAL if hl else (MUT85 if c_i == 0 else MUT78)
            set_run(p.add_run(), data[r_i][c_i], 13.5, color, bold=(c_i == 0 or hl))
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x18, 0x24, 0x36) if c_i == 0 else RGBColor(0x14, 0x1F, 0x30)
tb(s, 830, 152, 386, 130, [[("先看\u201c形态\u201d再看\u201c跑分\u201d：有没有 CLI 决定了它能否进你的 CI/CD 和自动化流水线——这是跑分榜不会告诉你的事。", dict(size=17, color=MUT72))]], line_spacing=1.7)
rect(s, 830, 300, 4, 104, fill=BLUE)
tb(s, 848, 300, 368, 108, [[("生态绑定是双刃剑：绑得深，日常集成省力；但换云、换办公套件时迁移成本也高。", dict(size=17, color=MUT72))]], line_spacing=1.7)
judge_card(s, 830, 470, 386, 156, "跑分是必要条件而非充分条件：SWE-bench 72.3% 只说明\u201c能做对题\u201d，能不能做对你的题，要靠下一页的 PoC。", 17)
footer(s, "格局篇", "09 / 21")

# ---------- P10 ----------
s = new_slide(prs)
header(s, "选型决策：按生态位匹配，不按跑分", "四个入口，对应四种真实处境——先定位自己在哪条岔路上")
rect(s, 64, 206, 150, 74, fill=TEAL_BG, line=TEAL_BR, radius=10)
tb(s, 64, 222, 150, 28, [[("你的场景", dict(size=19, bold=True))]], align=PP_ALIGN.CENTER)
tb(s, 64, 252, 150, 20, [[("START HERE", dict(size=12, color=MUT55, font=MONO))]], align=PP_ALIGN.CENTER)
routes = [
    (158, "微信 / 小程序 / 企微生态", "做微信端产品 → 腾讯 CodeBuddy：文档、组件库、云开发原生打通", True),
    (226, "阿里云栈 / 金融政务高合规", "要私有化与长上下文 → 阿里 Qoder：企业版私有化，100 万上下文", False),
    (294, "个人学习 / 学生 / 快速 MVP", "零成本起步 → 字节 Trae：全免费、国内直连、SOLO 一句话出项目", False),
    (362, "中文复杂业务 / 强数据安全", "要过程可回溯 → 百度文心快码：SPEC 模式 + 代码库级 RAG + 安全扫描", False),
]
for y, t1, t2, hl in routes:
    arrow(s, 214, 243, 464, y + 27, TEAL if hl else BLUE, 1.75)
    rect(s, 464, y, 752, 54, fill=CARD, line=TEAL_BR if hl else BLUE_BR, radius=9)
    tb(s, 488, y + 7, 700, 24, [[(t1, dict(size=16, bold=True, color=TEAL if hl else BLUE))]])
    tb(s, 488, y + 31, 710, 20, [[(t2, dict(size=13, color=MUT65))]])
rect(s, 64, 442, 1152, 178, fill=CARD, line=CARD_BR, radius=12)
tb(s, 92, 458, 400, 26, [[("选型三原则", dict(size=19, bold=True))]])
prins = [
    ("1", "生态优先", "你在哪个云、哪个办公套件里，就优先看谁——集成成本比功能差距更真实。"),
    ("2", "形态次之", "插件、独立 IDE 还是 CLI？决定的是工作流，而工作流决定习惯能否留存。"),
    ("3", "PoC 定案", "拿自己最复杂的三个真实任务跑两周，采纳率和返工率自己说话。"),
]
for i, (no, t, d) in enumerate(prins):
    x = 92 + i * 380
    if i:
        vline(s, x - 24, 500, 96, CARD_BR, 1)
    tb(s, x, 496, 44, 44, [[(no, dict(size=34, bold=True, color=TEAL, font=MONO))]])
    tb(s, x + 52, 494, 300, 26, [[(t, dict(size=18, bold=True))]])
    tb(s, x + 52, 524, 300, 76, [[(d, dict(size=14.5, color=MUT65))]], line_spacing=1.55)
footer(s, "格局篇", "10 / 21")

# ---------- P11 ----------
section_slide("03", os.path.join(A, "ch_practice.png"), "实践篇", "把默契写进仓库，而不是每次重讲",
              "上下文工程、规则文件、任务简报、增量循环、自主度分级、安全护栏——本章是全课重点，也是你带得走的部分。")

# ---------- P12 ----------
s = new_slide(prs)
header(s, "上下文工程：六层放对位置", "Agent 变笨的头号原因：什么都往提示词里塞", TEAL)
layers = [
    ("L1", "任务契约 Prompt", "只放四要素：目标、上下文指针、约束、完成标准", "100%", True),
    ("L2", "规则文件 AGENTS.md", "架构、命令、约定、边界——写一次，永久生效", "93%", True),
    ("L3", "Skills 可复用工作流", "重复 3 次以上的多步流程，封装成技能", "86%", False),
    ("L4", "MCP 外部工具", "日志、数据库、云服务——按需接入，控权限", "79%", False),
    ("L5", "Subagents 子智能体", "独立上下文的分身，隔离噪音、并行干活", "72%", False),
    ("L6", "Artifacts 产物沉淀", "生成的代码、文档、报告——沉淀回仓库", "65%", False),
]
for i, (no, name, what, w, hl) in enumerate(layers):
    y = 152 + i * 80
    ww = int(700 * int(w[:-1]) / 100)
    rect(s, 64, y, ww, 70, fill=TEAL_BG if hl else CARD, line=TEAL_BR if hl else CARD_BR, radius=9)
    tb(s, 84, y + 20, 60, 30, [[(no, dict(size=21, bold=True, color=TEAL if hl else BLUE, font=MONO))]])
    tb(s, 148, y + 22, 210, 28, [[(name, dict(size=18, bold=True))]])
    tb(s, 366, y + 12, ww - 320, 50, [[(what, dict(size=14.5, color=MUT68))]], line_spacing=1.4, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 808, 152, 408, 120, [[("六层是一个分流漏斗：静态、稳定的信息往下沉（规则、技能、工具），动态、一次性的信息往上浮（本次任务的契约）。", dict(size=19, color=TXT))]], line_spacing=1.7)
tb(s, 808, 292, 408, 110, [[("判断一条信息放哪一层的标准很简单——", dict(size=17, color=MUT65)), ("换个任务它还有用吗？", dict(size=17, bold=True, color=TEAL)), ("有用就沉下去，没用就留在 L1。", dict(size=17, color=MUT65))]], line_spacing=1.7)
judge_card(s, 808, 470, 408, 160, "上下文不是越多越好：无关信息会稀释关键规则。主线程越干净，Agent 越聪明。", 17.5)
footer(s, "实践篇", "12 / 21")

# ---------- P13 ----------
s = new_slide(prs)
header(s, "规则文件：AGENTS.md 怎么写", "60,000+ 仓库采用、30+ 工具原生读取 · 已由 Linux 基金会托管维护")
code_lines = [
    ("# AGENTS.md — 项目约定（手写，< 150 行）", TEAL),
    ("", TXT),
    ("## 技术栈", TEAL),
    ("- Next.js 15 App Router / TypeScript / Tailwind", TXT),
    ("- 包管理器：pnpm（禁止混用 npm / yarn）", TXT),
    ("", TXT),
    ("## 常用命令", TEAL),
    ("- 测试：pnpm test   类型检查：pnpm typecheck", TXT),
    ("- 改完业务代码必须跑测试，失败要解释原因", TXT),
    ("", TXT),
    ("## 约定（一条约定，配一个示例）", TEAL),
    ("- 新组件一律函数式 + hooks，禁止 class 组件", TXT),
    ("- 错误处理：让异常上抛，由全局 handler 统一捕获", TXT),
    ("  // 正确：函数内不 try/catch", RGBColor(0x6A, 0x99, 0x55)),
    ("  // 错误：局部 catch 后只打日志", RGBColor(0x6A, 0x99, 0x55)),
    ("", TXT),
    ("## 边界（红线）", TEAL),
    ("- 禁止修改 src/generated/**（自动生成目录）", TXT),
    ("- 禁止硬编码密钥、Token", TXT),
    ("- 新增外部依赖必须说明理由", TXT),
]
rect(s, 64, 148, 668, 470, fill=CODE_BG, line=CARD_BR, radius=10)
rect(s, 64, 148, 668, 34, fill=RGBColor(0x16, 0x22, 0x36), radius=10)
for i, c in enumerate((0xFF5F57, 0xFEBC2E, 0x28C840)):
    rect(s, 84 + i * 22, 159, 12, 12, fill=RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255), radius=6)
paras = [[(t, dict(size=13.5, color=c, font=MONO))] for t, c in code_lines]
tb(s, 92, 200, 620, 400, paras, line_spacing=1.32)
tips = [
    ("手写，不要生成", "实测 LLM 生成的 AGENTS.md 在 8 组对照里 5 组降低任务成功率，平均多走 2.45–3.92 步。"),
    ("150 行以内", "超过 150 行收益递减，推理成本反升 20–23%；上限 32KiB，短才有效。"),
    ("一条约定配一个示例", "一个反例代码胜过三段描述：\u201cShow, don't tell\u201d 是唯一有效的写法。"),
    ("随 PR 一起更新", "规则变了当场改；过时的规则比没有规则更糟——它在主动误导 Agent。"),
]
for i, (t, d) in enumerate(tips):
    y = 148 + i * 94
    rect(s, 776, y + 4, 26, 26, fill=BLUE_BG, line=BLUE_BR, radius=13)
    tb(s, 776, y + 9, 26, 18, [[(str(i + 1), dict(size=13, color=BLUE, font=MONO))]], align=PP_ALIGN.CENTER)
    tb(s, 818, y, 420, 26, [[(t, dict(size=18, bold=True))]])
    tb(s, 818, y + 30, 420, 56, [[(d, dict(size=14.5, color=MUT65))]], line_spacing=1.5)
rect(s, 776, 528, 440, 90, fill=TEAL_BG, line=TEAL_BR, radius=10)
tb(s, 796, 540, 400, 70, [[("本课判断：", dict(size=16.5, bold=True, color=TEAL)), ("空话会挤掉真规则。\u201c写干净的、可维护的代码\u201d这种话现在就删掉——Agent 本来就会。", dict(size=16.5, color=TXT))]], line_spacing=1.6)
footer(s, "实践篇", "13 / 21")

# ---------- P14 ----------
s = new_slide(prs)
header(s, "任务简报四要素：30 秒定生死", "产出的差距，八成在开工前 30 秒就决定了")
rect(s, 64, 148, 340, 490, fill=TEAL_BG, line=TEAL_BR, radius=12)
tb(s, 90, 174, 290, 70, [[("每条只写", dict(size=24, bold=True))], [("不可省略的信息", dict(size=24, bold=True))]], line_spacing=1.4)
elems = [
    ("Goal", "要解决什么问题，一句话"),
    ("Context", "相关代码与日志的路径指针"),
    ("Constraints", "不许动什么、不许加什么"),
    ("Done-when", "什么条件算完成、如何验证"),
]
for i, (k, d) in enumerate(elems):
    y = 268 + i * 66
    rect(s, 90, y, 6, 44, fill=TEAL, radius=3)
    tb(s, 110, y - 2, 280, 24, [[(k, dict(size=17, bold=True, color=TEAL, font=MONO))]])
    tb(s, 110, y + 22, 280, 24, [[(d, dict(size=14.5, color=MUT68))]])
tb(s, 90, 544, 290, 70, [[("静态约束写在 AGENTS.md，这里只放本次任务的增量。", dict(size=14.5, color=MUT55))]], line_spacing=1.6)
rect(s, 440, 148, 776, 296, fill=CARD, line=CARD_BR, radius=12)
tb(s, 468, 166, 500, 26, [[("一份合格简报的样子", dict(size=19, bold=True))]])
brief = [
    ("Goal", "修复优惠券接口返回 500 时订单无法提交的问题"),
    ("Context", "相关代码 apps/web/src/checkout；错误日志 logs/checkout-500.log"),
    ("Constraints", "不改支付 SDK；不新增依赖；报错文案风格与现有保持一致"),
    ("Done-when", "新增测试通过；本地下单全链路自测通过；根因写入 PR 描述"),
]
for i, (k, v) in enumerate(brief):
    y = 206 + i * 58
    rect(s, 468, y, 128, 32, fill=TEAL_BG2, radius=6)
    tb(s, 468, y + 6, 128, 22, [[(k, dict(size=14.5, bold=True, color=TEAL, font=MONO))]], align=PP_ALIGN.CENTER)
    tb(s, 616, y + 2, 580, 30, [[(v, dict(size=15.5, color=MUT85))]], line_spacing=1.4)
rect(s, 440, 466, 379, 172, fill=CARD, line=CARD_BR, radius=12)
tb(s, 462, 482, 200, 22, [[("✗ 坏简报", dict(size=15, bold=True, color=MUT50))]])
tb(s, 462, 510, 335, 26, [[("\u201c帮我修一下下单报错\u201d", dict(size=16.5, color=MUT78, font=MONO))]])
tb(s, 462, 544, 335, 80, [[("Agent 只能靠猜：猜范围、猜边界、猜\u201c修好\u201d的标准——猜错三次，信任归零。", dict(size=14.5, color=MUT55))]], line_spacing=1.55)
rect(s, 837, 466, 379, 172, fill=TEAL_BG, line=TEAL_BR, radius=12)
tb(s, 859, 482, 200, 22, [[("✓ 好简报的回报", dict(size=15, bold=True, color=TEAL))]])
tb(s, 859, 510, 335, 52, [[("第一次产出就接近可用，返工轮次从 3–4 轮降到 1 轮。", dict(size=16, color=TXT))]], line_spacing=1.55)
tb(s, 859, 572, 335, 60, [[("本课判断：", dict(size=14.5, bold=True, color=TEAL)), ("加约束比加形容词有效——\u201c认真写\u201d没有用，\u201c不改支付 SDK\u201d才有用。", dict(size=14.5, color=MUT68))]], line_spacing=1.55)
footer(s, "实践篇", "14 / 21")

# ---------- P15 ----------
s = new_slide(prs)
header(s, "增量交付循环：小步、验证、可回滚", "粒度决定回滚成本：一次一个切片，别让 Agent 一口气跑完整个需求")
nodes = [
    (86, 174, "① 拆", "需求切成可验证的切片", True),
    (852, 174, "② 派", "一次只派一个，带四要素", False),
    (852, 288, "③ 验", "测试 + 人审 diff，不过就回滚", False),
    (86, 288, "④ 收", "立即 commit，沉淀上下文", False),
]
arrow(s, 290, 206, 848, 206)
arrow(s, 952, 240, 952, 284)
arrow(s, 848, 320, 292, 320)
arrow(s, 186, 284, 186, 240)
for x, y, t, d, hl in nodes:
    rect(s, x, y, 200, 64, fill=TEAL_BG if hl else CARD, line=TEAL_BR if hl else BLUE_BR, radius=10)
    tb(s, x, y + 8, 200, 26, [[(t, dict(size=19, bold=True, color=TEAL if hl else BLUE))]], align=PP_ALIGN.CENTER)
    tb(s, x, y + 36, 200, 20, [[(d, dict(size=12.5, color=MUT65))]], align=PP_ALIGN.CENTER)
tb(s, 500, 232, 280, 26, [[("每圈 ≤ 1 天", dict(size=17, bold=True, color=MUT85))]], align=PP_ALIGN.CENTER)
tb(s, 500, 262, 280, 20, [[("SMALL BATCH, FAST FEEDBACK", dict(size=12, color=MUT50, font=MONO))]], align=PP_ALIGN.CENTER)
rect(s, 64, 420, 700, 230, fill=CARD, line=CARD_BR, radius=12)
tb(s, 92, 436, 400, 26, [[("每一步的硬检查点", dict(size=18, bold=True))]])
checks = [
    ("拆", "这个切片能独立验收吗？说不清\u201c怎么验\u201d就还没拆好。"),
    ("派", "四要素齐了吗？缺 Constraints 的简报退回重写。"),
    ("验", "测试红了吗？diff 逐行看过了吗？没看就等于没验。"),
    ("收", "commit message 写清\u201c为什么\u201d了吗？它是下一轮的上下文。"),
]
for i, (k, q) in enumerate(checks):
    x = 92 + (i % 2) * 336
    y = 476 + (i // 2) * 82
    rect(s, x, y, 30, 30, fill=TEAL_BG2, line=TEAL_BR, radius=15)
    tb(s, x, y + 5, 30, 20, [[(k, dict(size=14, color=TEAL, font=MONO))]], align=PP_ALIGN.CENTER)
    tb(s, x + 40, y - 2, 272, 78, [[(q, dict(size=14.5, color=MUT68))]], line_spacing=1.5)
rect(s, 788, 420, 428, 230, fill=TEAL_BG, line=TEAL_BR, radius=12)
rect(s, 812, 446, 10, 10, fill=TEAL, radius=5)
tb(s, 832, 440, 200, 22, [[("本课判断", dict(size=16, bold=True, color=TEAL))]])
tb(s, 812, 472, 380, 80, [[("一句\u201c帮我把整个模块写完\u201d得到的代码，你不敢上线；四个检查点走过的切片，你敢。", dict(size=17, color=TXT))]], line_spacing=1.7)
tb(s, 812, 566, 380, 70, [[("循环的意义不是慢，恰恰是快——因为返工被限制在了一个切片之内。", dict(size=15, color=MUT68))]], line_spacing=1.6)
footer(s, "实践篇", "15 / 21")

# ---------- P16 ----------
s = new_slide(prs)
header(s, "自主度分级：把介入点画在事前", "敢放手的前提，是知道每一级该在哪里伸手")
grades = [
    ("L0", "补全级", "人写，AI 补全单行", "介入：按键即审，零额外成本", "52%", False),
    ("L1", "单文件级", "AI 修改单个文件", "介入：逐行审 diff，看测试", "66%", False),
    ("L2", "跨文件级", "AI 完成跨模块小特性", "介入：先审方案设计，再验收测结果", "82%", False),
    ("L3", "长周期自主级", "连续多日自主运行（实测曾有 16 天 / 265 次提交）", "介入：事前定边界与红线，里程碑验收兜底", "100%", True),
]
for i, (no, name, desc, check, w, hl) in enumerate(grades):
    y = 168 + i * 118
    ww = int(700 * int(w[:-1]) / 100)
    rect(s, 64, y, ww, 106, fill=TEAL_BG if hl else CARD, line=TEAL_BR if hl else CARD_BR, radius=10)
    tb(s, 86, y + 18, 60, 28, [[(no, dict(size=19, bold=True, color=TEAL if hl else BLUE, font=MONO))]])
    tb(s, 148, y + 18, 170, 28, [[(name, dict(size=19, bold=True))]])
    tb(s, 86, y + 54, ww - 44, 44, [[(check, dict(size=14.5, color=TEAL if hl else MUT65))]])
    tb(s, 330, y + 20, ww - 290, 24, [[(desc, dict(size=13.5, color=MUT50))]])
tb(s, 808, 152, 408, 130, [[("等级每升一级，人的角色就换一次：从\u201c写字的\u201d变成\u201c改稿的\u201d，再变成\u201c定标准、验结果的人\u201d。", dict(size=19, color=TXT))]], line_spacing=1.7)
rect(s, 808, 300, 408, 110, fill=CARD, line=CARD_BR, radius=10)
tb(s, 828, 312, 370, 24, [[("这些事始终不建议交给 Agent 自主决策", dict(size=14.5, bold=True, color=BLUE))]])
tb(s, 828, 340, 370, 62, [[("核心金融计算与估值逻辑、加密与鉴权实现、大规模遗留系统重构——AI 只做展示与辅助，结论必须由人写。", dict(size=14.5, color=MUT68))]], line_spacing=1.65)
rect(s, 808, 470, 408, 160, fill=TEAL_BG, line=TEAL_BR, radius=10)
tb(s, 828, 486, 370, 130, [[("本课判断：", dict(size=16.5, bold=True, color=TEAL)), ("Qoder 的 16 天自主运行实验证明了上限，也证明了结论——长周期自主仍需人类兜底验收。放手不放眼。", dict(size=16.5, color=TXT))]], line_spacing=1.65)
footer(s, "实践篇", "16 / 21")

# ---------- P17 ----------
s = new_slide(prs)
header(s, "企业落地前，先装好安全护栏", "无论选哪家工具，进研发环境前这四件事一件都不能少", TEAL)
rect(s, 30, 540, 340, 340, fill=RGBColor(0x16, 0x2A, 0x3A), radius=170)
tb(s, 64, 160, 380, 150, [[("4", dict(size=150, bold=True, color=TEAL, font=MONO))]])
tb(s, 220, 268, 200, 40, [[("件套", dict(size=30, bold=True))]])
tb(s, 64, 330, 380, 120, [[("Agent 的每一点效率提升，都来自它比你多拿到的权限；对应地，每一条权限都是一条风险敞口。", dict(size=18, color=MUT72))]], line_spacing=1.7)
tb(s, 64, 470, 380, 24, [[("TRUST, BUT VERIFY — AND LOG", dict(size=14, color=MUT50, font=MONO, spacing=1))]])
guards = [
    ("01", "沙箱隔离", "Agent 在受限容器中运行：非 root、根目录只读、默认禁公网"),
    ("02", "最小权限", "只授予\u201c读仓库 + 写工作区\u201d；shell 全权、删库、外发一律不给"),
    ("03", "调用审计", "每一次工具调用留痕：谁、何时、动了什么——可回溯、可问责"),
    ("04", "人工审批闸", "删库、资金、数据外发等高敏感动作零自动化，必须人工确认"),
]
for i, (no, name, d) in enumerate(guards):
    y = 152 + i * 106
    rect(s, 520, y, 696, 94, fill=CARD, line=CARD_BR, radius=10)
    rect(s, 520, y, 4, 94, fill=TEAL)
    tb(s, 548, y + 28, 60, 36, [[(no, dict(size=25, bold=True, color=TEAL, font=MONO))]])
    tb(s, 622, y + 30, 170, 32, [[(name, dict(size=21, bold=True))]])
    tb(s, 806, y + 18, 390, 60, [[(d, dict(size=15.5, color=MUT68))]], line_spacing=1.5, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 64, 596, 1152, 66, fill=TEAL_BG, line=TEAL_BR, radius=10)
rect(s, 92, 620, 10, 10, fill=TEAL, radius=5)
tb(s, 114, 608, 1080, 40, [[("本课判断：", dict(size=17.5, bold=True, color=TEAL)), ("先画笼子，再放 Agent。护栏建设的一次性投入，远低于一次数据泄露的代价。", dict(size=17.5, color=TXT))]])
footer(s, "实践篇", "17 / 21")

# ---------- P18 ----------
section_slide("04", os.path.join(A, "ch_rollout.png"), "落地篇", "从个人提效到团队可复制",
              "个人用得好不难，难的是让 20 个人用得一样好——这一章给出 90 天路线图和度量方法。")

# ---------- P19 ----------
s = new_slide(prs)
header(s, "团队落地：90 天路线图", "每个阶段都有明确交付物——没有交付物的时间点就是走形式")
masked(os.path.join(A, "scene_pair.png"), os.path.join(A, "p_pair.png"), 0, 0, bottom=(0.32, 235))
rect(s, 64, 148, 210, 490, fill=TEAL_BG, line=TEAL_BR, radius=12)
tb(s, 88, 174, 170, 130, [[("30", dict(size=38, bold=True, color=TEAL, font=MONO))], [("60", dict(size=38, bold=True, color=TEAL, font=MONO))], [("90", dict(size=38, bold=True, color=TEAL, font=MONO))]], line_spacing=1.15)
tb(s, 88, 322, 170, 28, [[("三步走", dict(size=18, bold=True))]])
tb(s, 88, 356, 164, 80, [[("先建规范、再跑试点、最后进流水线。", dict(size=14.5, color=MUT55))]], line_spacing=1.6)
tb(s, 88, 500, 164, 120, [[("一汽的路径印证了顺序：规范先行，90% 研发启用后，AI 代码占比 25%、整体效率 +14%。", dict(size=13.5, color=MUT50))]], line_spacing=1.6)
phases = [
    ("Day 1-30", "建规范", ["手写 AGENTS.md：命令、约定、边界红线", "全员对齐四要素简报与增量循环"], "交付物：仓库规则文件 v1"),
    ("Day 31-60", "跑试点", ["2-3 名骨干 × 3 类真实任务", "每周复盘：采纳率、返工率基线"], "交付物：试点数据基线"),
    ("Day 61-90", "进流水线", ["验证步骤固化进 CI，Agent 直接消费", "安全护栏四件套上线，推广全团队"], "交付物：团队手册 + 治理基线"),
]
for i, (d, t, items, out) in enumerate(phases):
    y = 148 + i * 168
    rect(s, 296, y, 446, 156, fill=CARD, line=CARD_BR, radius=10)
    tb(s, 318, y + 14, 120, 24, [[(d, dict(size=15, color=BLUE, font=MONO))]])
    tb(s, 440, y + 10, 200, 28, [[(t, dict(size=18, bold=True))]])
    for j, it in enumerate(items):
        tb(s, 318, y + 46 + j * 26, 410, 24, [[("· " + it, dict(size=13.5, color=MUT68))]])
    tb(s, 318, y + 122, 410, 22, [[(out, dict(size=13.5, color=TEAL))]])
pic(s, os.path.join(A, "p_pair.png"), 764, 148, 452, 490)
rect(s, 764, 560, 452, 78, fill=BG)
tb(s, 784, 616, 420, 22, [[("人机各占一个工位：人定边界，Agent 干活", dict(size=14, color=MUT78))]])
footer(s, "落地篇", "19 / 21")

# ---------- P20 ----------
s = new_slide(prs)
header(s, "度量什么，就得到什么", "四个该看的指标，五个最常踩的坑")
metrics = [
    ("采纳率", "看产出质量", "AI 产出被保留的比例。行业参照：建信金科代码智能生成采纳率超 30%。"),
    ("返工率", "看过程成本", "AI 产出被打回重做的平均轮次，直接反映简报与规则文件的质量。"),
    ("缺陷密度", "看质量底线", "AI 参与代码的千行缺陷数，对照纯人工基线——守住质量红线。"),
    ("端到端时长", "看业务结果", "需求到上线的周期变化，这才是老板视角的真实收益。"),
]
for i, (k, tag, d) in enumerate(metrics):
    y = 152 + i * 120
    rect(s, 64, y, 620, 108, fill=CARD, line=CARD_BR, radius=10)
    rect(s, 64, y, 4, 108, fill=BLUE)
    tb(s, 92, y + 16, 240, 28, [[(k, dict(size=19, bold=True))]])
    tb(s, 320, y + 20, 200, 22, [[(tag, dict(size=13.5, color=TEAL))]])
    tb(s, 92, y + 50, 566, 48, [[(d, dict(size=14.5, color=MUT65))]], line_spacing=1.5)
tb(s, 728, 152, 488, 28, [[("五大高频误区", dict(size=19, bold=True))]])
pits = [
    ("统计\u201c生成了多少行\u201d", "奖励注水行为——改看采纳后存活代码。"),
    ("买工具不建规范", "没有 AGENTS.md，20 个人 20 种用法，产出互不相认。"),
    ("让 Agent 拍架构决策", "方案必须人审，Agent 交给你的默认是\u201c最平均的答案\u201d。"),
    ("一次派整个模块", "回滚成本爆炸，切回增量循环。"),
    ("把输出当终稿", "测试、评审、审计一样不能省——它写得快，不代表你可以审得快。"),
]
for i, (t, d) in enumerate(pits):
    y = 190 + i * 66
    rect(s, 728, y, 488, 58, fill=RGBColor(0x16, 0x21, 0x32), line=CARD_BR, radius=9)
    tb(s, 748, y + 16, 30, 26, [[(str(i + 1), dict(size=18, bold=True, color=BLUE, font=MONO))]])
    tb(s, 786, y + 8, 220, 22, [[(t, dict(size=15.5, bold=True))]])
    tb(s, 786, y + 32, 414, 22, [[(d, dict(size=13, color=MUT55))]])
rect(s, 728, 528, 488, 110, fill=TEAL_BG, line=TEAL_BR, radius=10)
tb(s, 748, 544, 448, 80, [[("本课判断：", dict(size=16, bold=True, color=TEAL)), ("度量是行为的指挥棒——指标错了，团队会精准地跑向错误方向。", dict(size=16, color=TXT))]], line_spacing=1.6)
footer(s, "落地篇", "20 / 21")

# ---------- P21 ----------
s = new_slide(prs)
rect(s, 0, 0, 1280, 6, fill=TEAL)
rect(s, -100, -140, 480, 480, fill=RGBColor(0x16, 0x2A, 0x3A), radius=240)
rect(s, 1000, 400, 520, 520, fill=RGBColor(0x15, 0x25, 0x3A), radius=260)
tb(s, 160, 150, 960, 80, [[("\u201c", dict(size=110, bold=True, color=RGBColor(0x2A, 0x5A, 0x5F), font=MONO))]], align=PP_ALIGN.CENTER)
tb(s, 160, 240, 960, 150, [[("决定产出的不是模型，", dict(size=44, bold=True))], [("而是", dict(size=44, bold=True)), ("你给它的上下文与边界", dict(size=44, bold=True, color=TEAL)), ("。", dict(size=44, bold=True))]], align=PP_ALIGN.CENTER, line_spacing=1.5)
tb(s, 160, 420, 960, 34, [[("把默契写进仓库，Agent 才会成为团队的一员。", dict(size=19, color=MUT65))]], align=PP_ALIGN.CENTER)
acts = [
    ("01", "本周内为自己的仓库建一个 AGENTS.md（手写，< 150 行）"),
    ("02", "用一个真实小需求跑通一轮\u201c拆-派-验-收\u201d增量循环"),
    ("03", "对照安全四件套自查权限，补上调用审计"),
]
for i, (no, act) in enumerate(acts):
    x = 160 + i * 324
    rect(s, x, 490, 308, 110, fill=CARD, line=CARD_BR, radius=10)
    tb(s, x + 20, 502, 80, 30, [[(no, dict(size=22, bold=True, color=TEAL, font=MONO))]])
    tb(s, x + 20, 536, 268, 56, [[(act, dict(size=15, color=MUT78))]], line_spacing=1.55)
tb(s, 160, 640, 600, 24, [[("Coding Agent 实战指南 · 研发效能内训 · 2026.10", dict(size=15, color=MUT50))]])
tb(s, 720, 640, 400, 24, [[("21 / 21", dict(size=15, color=MUT50, font=MONO))]], align=PP_ALIGN.RIGHT)

prs.save(OUT)
print("saved:", OUT, os.path.getsize(OUT), "bytes,", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
