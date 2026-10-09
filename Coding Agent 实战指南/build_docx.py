# -*- coding: utf-8 -*-
"""将《Coding Agent 工程实践手册》Markdown 渲染为排版良好的 Word 文档。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "Coding Agent 工程实践手册.md"
DST = "Coding Agent 工程实践手册.docx"

NAVY = RGBColor(0x0F, 0x1B, 0x2D)
TEAL = RGBColor(0x0E, 0x8F, 0x7E)
BLUE = RGBColor(0x2A, 0x5F, 0xC8)
GRAY = RGBColor(0x5A, 0x64, 0x72)
CODE_BG = "F4F6F9"
QUOTE_BG = "EEF5F3"
TH_BG = "E8EDF5"

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"


def set_run(run, size=10.5, bold=False, color=None, font=BODY_FONT, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), font)
    rf.set(qn('w:hAnsi'), font)
    rf.set(qn('w:eastAsia'), font if font != CODE_FONT else BODY_FONT)


def shade(el, color):
    pr = el.get_or_add_tcPr() if el.tag.endswith('tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    pr.append(shd)


def para_spacing(p, before=2, after=4, line=1.35):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`\n]+?`)')


def add_inline(p, text, size=10.5, color=None, font=BODY_FONT, bold=False):
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**'):
            set_run(p.add_run(seg[2:-2]), size, True, color, font)
        elif seg.startswith('`') and seg.endswith('`'):
            set_run(p.add_run(seg[1:-1]), size - 0.5, bold,
                    RGBColor(0xB0, 0x3A, 0x5A), CODE_FONT)
        else:
            set_run(p.add_run(seg), size, bold, color, font)


def add_code_block(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    shade(cell._tc, CODE_BG)
    cell.width = Cm(16.6)
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Pt(4)
        set_run(p.add_run(ln if ln else ' '), 8.5, False, RGBColor(0x22, 0x33, 0x44), CODE_FONT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def add_table(doc, rows):
    ncol = len(rows[0])
    tbl = doc.add_table(rows=len(rows), cols=ncol)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    total = Cm(16.6)
    w = total / ncol
    for r_i, row in enumerate(rows):
        for c_i in range(ncol):
            cell = tbl.cell(r_i, c_i)
            cell.width = w
            txt = row[c_i] if c_i < len(row) else ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            size = 8.5 if ncol >= 4 else 9.5
            if r_i == 0:
                shade(cell._tc, TH_BG)
                add_inline(p, txt, size, NAVY, BODY_FONT, True)
            else:
                add_inline(p, txt, size, RGBColor(0x2B, 0x33, 0x3D))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


BLOCK_START = re.compile(r'^(#{1,6}\s|\||```|>|\s*[-*]\s|\s*\d+\.\s|---)')


def continuation(lines, i):
    """收集软换行续行：以空白缩进开头、且不是新块首行的非空行。"""
    out = []
    n = len(lines)
    while i < n:
        raw = lines[i]
        s = raw.strip()
        if not s or BLOCK_START.match(raw):
            break
        if raw[:1] not in (' ', '\t', '　'):
            break
        out.append(s)
        i += 1
    return out, i


def build():
    src_text = open(SRC, encoding='utf-8').read()
    lines = src_text.split('\n')

    doc = Document()
    st = doc.styles['Normal']
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)

    for nm, sz, col in (('Heading 1', 20, NAVY), ('Heading 2', 15, NAVY),
                        ('Heading 3', 12.5, TEAL), ('Heading 4', 11, BLUE)):
        try:
            s = doc.styles[nm]
        except KeyError:
            continue
        s.font.name = BODY_FONT
        s.font.size = Pt(sz)
        s.font.color.rgb = col
        s.font.bold = True
        s.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)

    i = 0
    n = len(lines)
    title_done = False
    while i < n:
        ln = lines[i].rstrip()
        s = ln.strip()

        # fenced code
        if s.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i].rstrip('\n'))
                i += 1
            i += 1
            add_code_block(doc, buf)
            continue

        if not s:
            i += 1
            continue

        if s == '---':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue

        # table
        if s.startswith('|') and i + 1 < n and re.match(r'^\|[\s:\-\|]+\|$', lines[i + 1].strip()):
            rows = []
            header = [c.strip() for c in s.strip('|').split('|')]
            rows.append(header)
            i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            add_table(doc, rows)
            continue

        # headings
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            lvl = len(m.group(1))
            txt = m.group(2).strip()
            if lvl == 1 and not title_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para_spacing(p, 60, 6, 1.1)
                set_run(p.add_run(txt), 26, True, NAVY)
                title_done = True
            else:
                lvl = min(lvl, 4)
                p = doc.add_paragraph(style=f'Heading {lvl}')
                para_spacing(p, 12 if lvl <= 2 else 8, 5, 1.2)
                if lvl == 2:
                    set_run(p.add_run(txt), 15, True, NAVY)
                elif lvl == 3:
                    set_run(p.add_run(txt), 12.5, True, TEAL)
                else:
                    set_run(p.add_run(txt), 11, True, BLUE)
            i += 1
            continue

        # blockquote callout
        if s.startswith('>'):
            buf = []
            while i < n and lines[i].strip().startswith('>'):
                seg = lines[i].strip().lstrip('>').strip()
                if seg and seg != '—':
                    buf.append(seg)
                i += 1
            txt = ' '.join(buf)
            if txt.startswith('⚠️'):
                txt = txt.lstrip('⚠️').strip()
                txt = '⚠ ' + txt
            tbl = doc.add_table(rows=1, cols=1)
            cell = tbl.cell(0, 0)
            shade(cell._tc, QUOTE_BG)
            cell.width = Cm(16.6)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.3
            add_inline(p, txt, 10, RGBColor(0x1B, 0x3A, 0x36), BODY_FONT, True)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # ordered / unordered list
        mo = re.match(r'^(\d+)\.\s+(.*)$', s)
        mu = re.match(r'^[-*]\s+(.*)$', s)
        if mo or mu:
            items = []
            if mo:
                while i < n and re.match(r'^\d+\.\s+', lines[i].strip()):
                    it = re.sub(r'^\d+\.\s+', '', lines[i].strip())
                    i += 1
                    cont, i = continuation(lines, i)
                    items.append(it + ''.join(cont))
            else:
                while i < n and re.match(r'^[-*]\s+', lines[i].strip()):
                    it = re.sub(r'^[-*]\s+', '', lines[i].strip())
                    i += 1
                    cont, i = continuation(lines, i)
                    items.append(it + ''.join(cont))
            for idx, it in enumerate(items, 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.55)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.3
                marker = f'{idx}. ' if mo else '• '
                set_run(p.add_run(marker), 10, True, TEAL if mo else BLUE)
                add_inline(p, it, 10)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue

        # paragraph（吸收软换行续行）
        body = s
        i += 1
        cont, i = continuation(lines, i)
        if cont:
            body += ''.join(cont)
        p = doc.add_paragraph()
        para_spacing(p, 2, 4, 1.4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, body, 10.5)

    doc.save(DST)
    print('saved', DST, 'paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))


if __name__ == '__main__':
    build()
